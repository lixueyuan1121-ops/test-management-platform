"""Requirement review/vision integration, isolated DB and fake model only.

Run: .venv/bin/python -m unittest scripts.test_requirement_analysis -v
"""
import copy
import io
import json
import unittest
from datetime import date
from unittest.mock import patch

from fastapi.testclient import TestClient
from PIL import Image
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.deps import get_current_user
from app.core.enums import ProjectRole
from app.db.session import Base, get_db
from app.main import app
from app.models import (AiJob, Project, ProjectMember, RequirementBaseline,
                        RequirementSource, Task, User)
from app.schemas.requirement_analysis import RequirementDraft
from app.services import requirement_analysis as review, requirement_sources as sources
from app.services.generators import deepseek_runner


def png(width=200, height=100):
    output = io.BytesIO()
    Image.new("RGB", (width, height), "white").save(output, "PNG")
    return output.getvalue()


def draft():
    return RequirementDraft.model_validate({"summary": "文件删除确认", "scope": "Windows", "rules": [
        {"id": "R1", "title": "保护目录删除", "condition": "文件位于保护目录", "action": "点击删除",
         "expected": "显示确认弹窗", "source_type": "explicit", "source_quote": "保护目录删除必须确认",
         "criteria": [{"id": "R1-C1", "text": "保护目录删除显示确认弹窗"},
                      {"id": "R1-C2", "text": "取消确认后文件保留"}]},
        {"id": "R2", "title": "可选回收策略", "source_type": "inferred"}],
        "questions": [{"id": "Q1", "question": "是否启用回收站？", "rule_ids": ["R2"]}]
    }).model_dump()


class FakeEngine:
    is_available = staticmethod(lambda: True)

    def __init__(self, response):
        self.response = response
        self.calls = []

    def stream_generate(self, _text, **kwargs):
        self.calls.append({"prompt": kwargs["prompt_builder"](), "images": kwargs.get("images")})
        data = {"kind": "表格", "text": "|目录|行为|\n|保护目录|必须确认|", "uncertainties": ""} if kwargs.get("images") else self.response
        yield {"type": "result", "text": json.dumps(data, ensure_ascii=False)}


class ReviewAPITests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        Base.metadata.create_all(self.engine)
        self.Session = sessionmaker(bind=self.engine)
        self.db = self.Session()
        self.db.add_all([User(id=i, username=f"u{i}", name=f"u{i}", password_hash="x", is_platform_admin=False) for i in (1, 2, 3)])
        self.db.add_all([Project(id=1, name="P1", code="p1"), Project(id=2, name="P2", code="p2")])
        self.db.flush()
        self.db.add_all([ProjectMember(user_id=1, project_id=1, role=ProjectRole.member),
                         ProjectMember(user_id=2, project_id=2, role=ProjectRole.member),
                         ProjectMember(user_id=3, project_id=1, role=ProjectRole.guest),
                         Task(id=1, project_id=1, assigned_by=1, assigned_to=1, title="任务", assigned_date=date.today())])
        self.db.commit()
        self.uid = 1
        def get_test_db():
            with self.Session() as s:
                yield s
        app.dependency_overrides[get_db] = get_test_db
        app.dependency_overrides[get_current_user] = lambda: self.db.get(User, self.uid)
        self.client = TestClient(app)
        self.wake_patch = patch("app.services.ai_jobs.notify_new_job")
        self.wake_patch.start()
        self.fake = FakeEngine(draft())
        self.provider_patch = patch("app.services.generators.get_provider", return_value=self.fake)
        self.provider_patch.start()

    def tearDown(self):
        self.provider_patch.stop(); self.wake_patch.stop()
        app.dependency_overrides.clear()
        self.client.close(); self.db.close(); self.engine.dispose()

    def analyze(self, **kwargs):
        response = self.client.post("/api/ai/requirements/analyze", json={"project_id": 1, "task_id": 1,
                    "requirement": "保护目录删除必须确认", **kwargs})
        self.assertEqual(response.status_code, 200, response.text)
        result = response.json()["data"]
        self.db.expire_all()
        review.run_analysis_job(self.db, self.db.get(AiJob, result["job_id"]))
        return self.detail(result["analysis_id"])

    def detail(self, aid):
        return self.client.get(f"/api/ai/requirements/analyses/{aid}").json()["data"]

    def save(self, analysis):
        response = self.client.patch(f"/api/ai/requirements/analyses/{analysis['id']}", json={"revision": analysis["revision"], "draft": analysis["draft"]})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()["data"]

    def confirm(self, analysis, **kwargs):
        return self.client.post(f"/api/ai/requirements/analyses/{analysis['id']}/confirm", json={
            "revision": analysis["revision"], "source_hash": analysis["source_hash"], "scope_reviewed": True, **kwargs})

    def test_confirmation_generation_and_coverage(self):
        # Even malicious model output cannot impersonate human review.
        self.fake.response["rules"][0].update(status="confirmed", review_note="模型自行确认")
        self.fake.response["questions"][0]["answer"] = "模型自行决定"
        analysis = self.analyze()
        self.assertEqual(analysis["draft"]["rules"][0]["status"], "pending")
        self.assertEqual(analysis["draft"]["questions"][0]["answer"], "")
        self.assertEqual(self.confirm(analysis).status_code, 422)
        analysis["draft"]["rules"][0]["status"] = "confirmed"
        analysis = self.save(analysis)
        confirmation = self.confirm(analysis)
        self.assertEqual(confirmation.status_code, 200, confirmation.text)
        bid = confirmation.json()["data"]["baseline_id"]
        self.assertEqual(self.confirm(analysis).json()["data"]["baseline_id"], bid)
        self.assertEqual(self.confirm(analysis, scope_reviewed=False).status_code, 422)
        body = {"project_id": 1, "task_id": 1, "requirement": analysis["source_text"], "baseline_id": bid}
        self.assertEqual(self.client.post("/api/ai/testcases", json={**body, "baseline_id": None}).status_code, 400)
        self.assertEqual(self.client.post("/api/ai/testcases", json={**body, "requirement": "更改"}).status_code, 409)
        response = self.client.post("/api/ai/testcases", json=body)
        self.assertEqual(response.status_code, 200, response.text)
        gen = response.json()["data"]
        self.db.expire_all()
        job = self.db.get(AiJob, gen["job_id"])
        self.assertNotIn("requirement", json.loads(job.input))
        generated = {"cases": [{"title": "删除保护文件", "category": "功能", "priority": "P0", "steps": "点击删除", "expected": "显示确认弹窗", "criterion_ids": ["R1-C1"], "precondition": "准备保护文件"}], "raw": "中" * 60000, "errors": [], "meta": {}}
        from app.api.ai import run_testcase_gen_job
        with patch.object(review, "generate_from_baseline", return_value=generated):
            result = run_testcase_gen_job(self.db, job)
        self.assertNotIn("cases", result)  # Small queue pointers, full detail stays in domain tables.
        self.assertLess(len(json.dumps(result).encode()), 1000)
        cases = self.client.get(f"/api/ai/tasks/{gen['ai_task_id']}/cases").json()["data"]
        self.assertEqual(cases[0]["acceptance_links"][0]["criterion_id"], "R1-C1")
        self.assertEqual(cases[0]["precondition"], "准备保护文件")
        def coverage():
            return self.client.get(f"/api/ai/requirements/coverage/{gen['ai_task_id']}").json()["data"]
        cov = coverage()
        self.assertEqual((cov["total"], cov["linked"], cov["reviewed"]), (2, 1, 0))
        self.assertEqual(cov["pending_rules"], ["R2"])
        cid = cases[0]["id"]
        self.assertEqual(self.client.patch(f"/api/ai/testcases/{cid}", json={"review_status": "adopted"}).status_code, 200)
        self.assertEqual(coverage()["reviewed"], 1)
        self.client.patch(f"/api/ai/testcases/{cid}", json={"expected": "修改预期"})
        self.assertEqual(coverage()["reviewed"], 0)
        self.client.patch(f"/api/ai/testcases/{cid}", json={"review_status": "rejected"})
        self.assertEqual(coverage()["linked"], 0)
        analysis["draft"]["scope"] = "仅 Windows 11"
        updated = self.save(analysis)
        self.assertIsNone(updated["baseline_id"])
        self.assertTrue(coverage()["newer_draft"])
        self.assertEqual(self.client.post("/api/ai/testcases", json=body).status_code, 409)
        self.db.expire_all()
        self.assertEqual(json.loads(self.db.get(RequirementBaseline, bid).payload)["scope"], "Windows")
        self.assertEqual(self.confirm(analysis).status_code, 409)

    def test_questions_evidence_and_revision_conflicts(self):
        analysis = self.analyze()
        analysis["draft"]["rules"][0]["status"] = "confirmed"
        analysis["draft"]["questions"][0]["rule_ids"] = ["R1"]
        analysis = self.save(analysis)
        self.assertEqual(self.confirm(analysis).status_code, 422)
        analysis["draft"]["questions"][0]["answer"] = "暂不启用回收站"
        analysis["draft"]["rules"][0]["source_quote"] = "不存在的原文"
        analysis = self.save(analysis)
        self.assertEqual(analysis["draft"]["rules"][0]["source_type"], "inferred")
        self.assertEqual(self.confirm(analysis).status_code, 422)
        analysis["draft"]["rules"][0]["review_note"] = "产品确认补充此规则"
        analysis = self.save(analysis)
        self.assertEqual(self.confirm(analysis).status_code, 200)
        bad = copy.deepcopy(analysis["draft"]); bad["rules"][0]["source_material_ids"] = ["IMG404"]
        self.assertEqual(self.client.patch(f"/api/ai/requirements/analyses/{analysis['id']}", json={"revision": analysis["revision"], "draft": bad}).status_code, 422)
        deleted = copy.deepcopy(analysis["draft"]); deleted["questions"] = []
        self.assertEqual(self.client.patch(f"/api/ai/requirements/analyses/{analysis['id']}", json={"revision": analysis["revision"], "draft": deleted}).status_code, 422)

    def test_images_authorization_and_append_snapshot(self):
        first = self.client.post("/api/ai/extract-file", files={"file": ("matrix.png", png(), "image/png")}).json()["data"]
        sid = first["source_id"]
        self.assertNotIn("data", first["materials"][0])
        second = self.client.post(f"/api/ai/extract-file?append_to={sid}", files={"file": ("flow.png", png(), "image/png")}).json()["data"]
        self.assertEqual([m["id"] for m in second["materials"]], ["IMG1", "IMG2"])
        self.assertEqual(len(json.loads(self.db.get(RequirementSource, sid).materials)), 1)
        analysis = self.analyze(source_id=second["source_id"])
        self.assertEqual(len(analysis["visual_readings"]), 2)
        self.assertEqual(analysis["visual_readings"][0]["status"], "read")
        self.assertTrue(self.fake.calls[0]["images"][0]["data"])
        self.assertIn("|保护目录|必须确认|", self.fake.calls[-1]["prompt"])
        image_url = f"/api/ai/requirements/sources/{sid}/images/IMG1"
        response = self.client.get(image_url)
        self.assertEqual(response.headers["cache-control"], "private, no-store")
        self.assertEqual(Image.open(io.BytesIO(response.content)).size, (200, 100))
        self.uid = 2
        self.assertEqual(self.client.get(image_url).status_code, 403)
        self.assertEqual(self.client.get(f"/api/ai/requirements/analyses/{analysis['id']}").status_code, 403)
        self.uid = 3
        self.assertEqual(self.client.get(f"/api/ai/requirements/sources/{second['source_id']}/images/IMG1").status_code, 200)
        self.assertEqual(self.confirm(analysis).status_code, 403)

    def test_failed_images_and_truncated_source_never_silent(self):
        broken = self.client.post("/api/ai/extract-file", files={"file": ("bad.png", b"invalid", "image/png")}).json()["data"]
        analysis = self.analyze(source_id=broken["source_id"])
        self.assertEqual(analysis["visual_readings"][0]["status"], "failed")
        analysis["draft"]["rules"][0]["status"] = "confirmed"
        analysis = self.save(analysis)
        self.assertEqual(self.confirm(analysis).status_code, 422)
        self.assertEqual(self.confirm(analysis, confirmation_note="图片范围待澄清，仅确认正文的删除规则").status_code, 200)
        source = self.client.post("/api/ai/extract-file", files={"file": ("long.txt", ("中" * 60001).encode(), "text/plain")}).json()["data"]
        self.assertTrue(source["truncated"])
        self.assertEqual(len(source["text"]), 60000)
        self.assertTrue(source["warnings"])
        self.assertEqual(len(self.db.get(RequirementSource, source["source_id"]).text), 60001)


class ExtractionTests(unittest.TestCase):
    def test_word_table_images_and_pdf_scan(self):
        import docx
        doc = docx.Document()
        doc.add_paragraph("权限矩阵")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "模式"; table.cell(0, 1).text = "删除"
        table.cell(1, 0).text = "普通"; table.cell(1, 1).text = "确认"
        doc.add_picture(io.BytesIO(png()))
        output = io.BytesIO(); doc.save(output)
        text, materials, warnings = sources.from_file("test.docx", output.getvalue())
        self.assertIn("模式 | 删除\n普通 | 确认", text)
        self.assertEqual(len(materials), 1)
        scan = io.BytesIO(); Image.open(io.BytesIO(png())).save(scan, "PDF")
        text, images, warnings = sources.from_file("scan.pdf", scan.getvalue())
        self.assertTrue(images[0]["data"])
        self.assertIn("PDF 第 1 页", images[0]["location"])
        self.assertIn("第 1 页", text)

    def test_long_image_limits_and_feishu_pagination(self):
        materials = sources.Materials(); materials.add(png(200, 4000), "长图")
        self.assertEqual(len(materials.items), 3)
        self.assertIn("分片 3/3", materials.items[-1]["location"])
        with patch.object(sources, "MAX_IMAGES", 1):
            limited = sources.Materials(); limited.add(png(200, 4000), "长图")
            self.assertTrue(limited.warnings)
        pages = [ {"items": [{"block_id": "a", "image": {"token": "media1"}}], "has_more": True, "page_token": "p2"},
                  {"items": [{"block_id": "b", "image": {"token": "media2"}}], "has_more": False} ]
        with patch.object(sources.feishu, "_fetch_docx", return_value=("文档", "正文")), \
             patch.object(sources.feishu, "_api_get", side_effect=pages) as api, \
             patch.object(sources, "_download_media", side_effect=[png(), ValueError("无素材权限")]):
            title, text, images, warnings = sources.from_url("https://my.feishu.cn/docx/Abc123")
        self.assertEqual(len(images), 2)
        self.assertEqual(images[1]["error"], "无素材权限")
        self.assertEqual(api.call_args_list[1].args[1]["page_token"], "p2")
        with self.assertRaises(ValueError):
            sources._public_download("http://127.0.0.1/private")

    def test_vision_errors_and_openai_image_protocol(self):
        with patch.object(deepseek_runner.settings, "DEEPSEEK_VISION_MODEL", "configured-vision"):
            request = deepseek_runner._body("read", True, images=[{"mime_type": "image/png", "data": "cGlj"}])
        self.assertEqual(request["model"], "configured-vision")
        self.assertEqual(request["messages"][1]["content"][1]["image_url"]["url"], "data:image/png;base64,cGlj")
        class Failure:
            def stream_generate(self, *args, **kwargs):
                yield {"type": "error", "msg": "模型不支持图片"}
        result = review.read_image(Failure(), {"id": "IMG1", "data": "x", "mime_type": "image/png"})
        self.assertEqual(result["status"], "failed")
        self.assertIn("不支持图片", result["uncertainties"])
        with self.assertRaises(ValueError):
            review.parse_object('{"rules": [')

    def test_url_text_preserves_operators_and_html_table_columns(self):
        with patch.object(sources, "_public_download", return_value=(b"x < 5 and y > 2", "text/plain", "https://example.com/r")):
            self.assertEqual(sources.from_url("https://example.com/r")[1], "x < 5 and y > 2")
        page = '<html><table><tr><th>模式</th><th>删除</th></tr><tr><td>普通</td><td>确认</td></tr></table></html>'
        with patch.object(sources, "_public_download", return_value=(page.encode("gbk"), "text/html; charset=gbk", "https://example.com/r")):
            text = sources.from_url("https://example.com/r")[1]
        self.assertIn("模式 | 删除 |", text)
        self.assertIn("普通 | 确认 |", text)

    def test_generation_keeps_missing_and_invalid_references_visible(self):
        payload = draft(); payload["rules"][0]["status"] = "confirmed"
        fake = FakeEngine([])
        fake.parse_testcases = lambda *args, **kwargs: [{"title": "缺少依据", "steps": "操作", "expected": "结果", "criterion_ids": ["R999-C1"]}]
        with patch("app.api.ai._gen_once", return_value=("[]", {}, None)):
            result = review.generate_from_baseline(fake, payload, None, None, "")
        self.assertEqual(result["cases"][0]["criterion_ids"], [])
        self.assertTrue(any("未生成验收条件" in error for error in result["errors"]))


if __name__ == "__main__":
    unittest.main()
